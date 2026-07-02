from pathlib import Path
import re
import tempfile

from .types import DocumentMetadata, ParsedDocument, NOT_DETECTED


class DocumentParseError(Exception):
    pass


SUPPORTED_EXTENSIONS = {".pdf", ".doc", ".docx", ".txt"}


def parse_document(filename, content):
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise DocumentParseError("Unsupported file type. Please upload a PDF, DOC, DOCX, or TXT file.")
    if not content:
        raise DocumentParseError("This document appears to be empty.")

    if ext == ".txt":
        text = _decode_text(content)
    elif ext == ".docx":
        text = _extract_docx(content)
    elif ext == ".pdf":
        text = _extract_pdf(content)
    else:
        text = _extract_doc_best_effort(content)

    text = normalize_text(text)
    if not text:
        raise DocumentParseError("No readable text could be extracted from this document.")
    if len(text) < 80:
        raise DocumentParseError("This document has very little usable text. Try uploading a text-based PDF, DOCX, or TXT file.")

    return ParsedDocument(text=text, metadata=detect_metadata(filename, text))


def parse_extracted_text(filename, text):
    text = normalize_text(text)
    if not text:
        raise DocumentParseError("No readable text could be extracted from this document.")
    if len(text) < 80:
        raise DocumentParseError("This document has very little usable text. Try uploading a text-based PDF, DOCX, or TXT file.")
    return ParsedDocument(text=text, metadata=detect_metadata(filename, text))


def _decode_text(content):
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="ignore")


def _extract_docx(content):
    try:
        from docx import Document
    except Exception as exc:
        raise DocumentParseError(f"DOCX support is unavailable: {exc}") from exc

    with tempfile.NamedTemporaryFile(suffix=".docx") as tmp:
        tmp.write(content)
        tmp.flush()
        try:
            document = Document(tmp.name)
        except Exception as exc:
            raise DocumentParseError("This DOCX file is unreadable or password-protected.") from exc
        return "\n".join(p.text for p in document.paragraphs)


def _extract_pdf(content):
    try:
        import pdfplumber
    except Exception as exc:
        raise DocumentParseError(f"PDF support is unavailable: {exc}") from exc

    with tempfile.NamedTemporaryFile(suffix=".pdf") as tmp:
        tmp.write(content)
        tmp.flush()
        try:
            with pdfplumber.open(tmp.name) as pdf:
                return "\n".join(page.extract_text() or "" for page in pdf.pages)
        except Exception as exc:
            raise DocumentParseError("This PDF is unreadable, password-protected, scanned without OCR, or text extraction failed.") from exc


def _extract_doc_best_effort(content):
    decoded = _decode_text(content)
    readable = re.sub(r"[^\x09\x0A\x0D\x20-\x7E]+", " ", decoded)
    readable = normalize_text(readable)
    if len(readable) < 200:
        raise DocumentParseError("This legacy DOC file could not be reliably read. Please upload a DOCX, PDF, or TXT version if available.")
    return readable


def normalize_text(text):
    text = text.replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def detect_metadata(filename, text):
    return DocumentMetadata(
        uploaded_file_name=filename,
        document_title=_detect_title(text),
        event_or_meeting_date=_detect_date(text),
        jurisdiction=_detect_jurisdiction(text),
        agency_or_governing_body=_detect_agency(text),
        document_type=_detect_document_type(text),
    )


def _detect_title(text):
    for line in text.splitlines()[:20]:
        line = _clean_line(line)
        if not line or len(line) < 8 or len(line) > 140:
            continue
        if re.search(r"\b(page|table of contents|draft|confidential)\b", line, re.IGNORECASE):
            continue
        if re.search(r"\b(agenda|staff report|board packet|policy memo|ordinance|resolution|plan|report|notice)\b", line, re.IGNORECASE):
            return line
    return NOT_DETECTED


def _detect_date(text):
    patterns = [
        r"\b(?:January|February|March|April|May|June|July|August|September|October|November|December)\s+\d{1,2},\s+\d{4}\b",
        r"\b\d{1,2}/\d{1,2}/\d{2,4}\b",
        r"\b\d{4}-\d{2}-\d{2}\b",
    ]
    for pattern in patterns:
        match = re.search(pattern, text, flags=re.IGNORECASE)
        if match:
            return match.group(0)
    return NOT_DETECTED


def _detect_jurisdiction(text):
    patterns = [
        r"\b(City of [A-Z][A-Za-z .'-]+)",
        r"\b(County of [A-Z][A-Za-z .'-]+)",
        r"\b([A-Z][A-Za-z .'-]+ County)\b",
        r"\b([A-Z][A-Za-z .'-]+ School District)\b",
        r"\b([A-Z][A-Za-z .'-]+ Transit Authority)\b",
        r"\b(State of [A-Z][A-Za-z .'-]+)",
    ]
    for pattern in patterns:
        match = re.search(pattern, text)
        if match:
            return _clean_line(match.group(1))
    return NOT_DETECTED


def _detect_agency(text):
    for line in text.splitlines()[:60]:
        line = _clean_line(line)
        if re.search(r"\b(board|council|commission|committee|agency|authority|department|district)\b", line, re.IGNORECASE):
            if 6 <= len(line) <= 120:
                return line
    return NOT_DETECTED


def _detect_document_type(text):
    lowered = text.lower()
    candidates = [
        ("meeting agenda", "Meeting agenda"),
        ("board packet", "Board packet"),
        ("staff report", "Staff report"),
        ("policy memo", "Policy memo"),
        ("ordinance", "Ordinance"),
        ("resolution", "Resolution"),
        ("grant notice", "Grant notice"),
        ("planning document", "Planning document"),
        ("budget", "Budget document"),
        ("community engagement report", "Community engagement report"),
        ("public agency report", "Public agency report"),
    ]
    for needle, label in candidates:
        if needle in lowered:
            return label
    return NOT_DETECTED


def _clean_line(line):
    return re.sub(r"\s+", " ", line).strip(" -|")
